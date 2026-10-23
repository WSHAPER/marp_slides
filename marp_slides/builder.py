"""The native PPTX builder: a parsed deck to presentation bytes.

16:9 slide, Title Slide layout for title-only slides, Title and Content
for everything else; body placeholders for text flow, sequential native
shapes for code fences, tables and pictures; two-column text boxes for
column slides; chrome (footer/header/slide numbers) outside the content
flow; speaker notes as native notes slides.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterable, Mapping

import pptx.presentation
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.slide import Slide, SlideLayout
from pptx.text.text import TextFrame, _Paragraph
from pptx.util import Emu, Inches, Pt

from .blocks import (
    Block,
    BlockQuote,
    CodeBlock,
    ColumnBreak,
    ImageBlock,
    ListBlock,
    Paragraph,
    ParsedDeck,
    ParsedSlide,
    Run,
    Subheading,
    TableBlock,
)
from .elements import (
    add_chrome_box,
    add_code_box,
    add_picture,
    add_slide_number_field,
    add_table,
    add_text_block,
    apply_list_para,
    apply_quote_para,
    picture_geometry,
    quote_runs,
    style_subheading,
    write_runs,
)
from .geometry import (
    BLOCK_GAP,
    CHROME_RESERVE,
    COLUMN_GAP,
    COLUMN_WIDTH,
    COLUMN_WIDTH_PT,
    CONTENT_LEFT,
    CONTENT_WIDTH_PT,
    FLOW_BAND_IN,
    FLOW_BOTTOM,
    FLOW_TOP,
    FOOTER_TOP,
    HEADER_TOP,
    LIST_HEIGHT,
    MONO_FONT,
    PARA_HEIGHT,
    QUOTE_HEIGHT,
    SLIDE_HEIGHT_EMU,
    SLIDE_WIDTH_EMU,
    SUBHEAD_HEIGHT,
    TABLE_ROW_HEIGHT,
    fit_code,
)


def build_pptx(
    deck: ParsedDeck, presentation_title: str, images: Mapping[int, bytes] | None = None
) -> bytes:
    """Build the native, editable PPTX (16:9) from a parsed deck.

    ``images`` maps the deck's file-id references (``#file:N``, 1-based) to
    the downloaded picture bytes; it may be None/empty for decks without
    image references (the worker guarantees consistency via
    :func:`marp_slides.images.validate_image_references`)."""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(SLIDE_HEIGHT_EMU)
    prs.core_properties.title = presentation_title

    title_layout = _layout_by_name(prs, "Title Slide", 0)
    content_layout = _layout_by_name(prs, "Title and Content", 1)

    for slide in deck.slides:
        _build_slide(prs, slide, deck, title_layout, content_layout, images or {})

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _build_slide(
    prs: pptx.presentation.Presentation,
    slide: ParsedSlide,
    deck: ParsedDeck,
    title_layout: SlideLayout,
    content_layout: SlideLayout,
    images: Mapping[int, bytes],
) -> None:
    is_title_slide = not slide.blocks
    layout = title_layout if is_title_slide else content_layout
    s = prs.slides.add_slide(layout)
    write_runs(
        s.placeholders[0].text_frame.paragraphs[0], slide.title_runs
    )  # rich runs in the title placeholder (bold/links/code)

    background = slide.background or deck.background
    if background is not None:
        s.background.fill.solid()  # native <p:bg> solid fill
        s.background.fill.fore_color.rgb = RGBColor.from_string(  # type: ignore[no-untyped-call]
            background
        )

    # effective chrome: per-slide scoped overrides win over deck values
    footer_runs = (
        slide.footer_override if slide.footer_override is not None else deck.footer
    )
    header_runs = (
        slide.header_override if slide.header_override is not None else deck.header
    )

    _fill_slide_body(s, slide, is_title_slide, footer_runs, images)

    if footer_runs:
        add_chrome_box(s, footer_runs, FOOTER_TOP, "Deck Footer")
    if header_runs:
        add_chrome_box(s, header_runs, HEADER_TOP, "Deck Header")
    if deck.paginate:
        add_slide_number_field(s, slide.number)
    if slide.notes:
        s.notes_slide.notes_text_frame.text = "\n".join(slide.notes)


def _fill_slide_body(
    s: Slide,
    slide: ParsedSlide,
    is_title_slide: bool,
    footer_runs: tuple[Run, ...] | None,
    images: Mapping[int, bytes],
) -> None:
    """Pick the slide's content layout branch: placeholder body, two
    columns, sequential flow shapes, or nothing (title slide)."""
    if is_title_slide:
        _drop_placeholder(s, 1)  # unused subtitle
        return
    bottom = FLOW_BOTTOM - CHROME_RESERVE if footer_runs else FLOW_BOTTOM
    if slide.columns is not None:
        _drop_placeholder(s, 1)  # content renders as two column boxes
        _add_columns(s, slide, bottom)
        return
    if slide.needs_flow_layout:
        _drop_placeholder(s, 1)  # blocks render as sequential shapes
        flow_blocks(s, slide.blocks, bottom, images)
        return
    _fill_body(s.placeholders[1].text_frame, slide.blocks)


def _add_columns(slide_shape: Slide, slide: ParsedSlide, bottom: int) -> None:
    """Two-column layout: two side-by-side native text boxes. Content blocks
    before the ColumnBreak render in the left box, blocks after it in the
    right one; without a break everything stays in the left column."""
    left_blocks: list[Block] = []
    right_blocks: list[Block] = []
    target = left_blocks
    for block in slide.blocks:
        if isinstance(block, ColumnBreak):
            target = right_blocks
            continue
        target.append(block)
    height = int(bottom) - int(FLOW_TOP)
    for idx, blocks in enumerate((left_blocks, right_blocks)):
        left = Emu(int(CONTENT_LEFT) + idx * (int(COLUMN_WIDTH) + int(COLUMN_GAP)))
        box = slide_shape.shapes.add_textbox(
            left, Emu(int(FLOW_TOP)), COLUMN_WIDTH, Emu(height)
        )
        box.name = "Column 1" if idx == 0 else "Column 2"
        tf = box.text_frame
        tf.word_wrap = True
        if blocks:
            _fill_body(tf, blocks, width_pt=COLUMN_WIDTH_PT, avail_in=height / 914400)


def _layout_by_name(
    prs: pptx.presentation.Presentation, name: str, fallback_index: int
) -> SlideLayout:
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    return prs.slide_layouts[fallback_index]


def _drop_placeholder(slide: Slide, idx: int) -> None:
    ph = slide.placeholders[idx]
    ph._element.getparent().remove(ph._element)


def _fill_body(
    text_frame: TextFrame,
    blocks: Iterable[Block],
    width_pt: float = CONTENT_WIDTH_PT,
    avail_in: float = 5.6,
) -> None:
    """Write blocks into a text frame (body placeholder or column box),
    one native paragraph per block item."""
    first = True

    def next_para() -> _Paragraph:
        nonlocal first
        para = text_frame.paragraphs[0] if first else text_frame.add_paragraph()  # type: ignore[no-untyped-call]
        first = False
        return para

    for block in blocks:
        if isinstance(block, Paragraph):
            write_runs(next_para(), block.runs)
        elif isinstance(block, Subheading):
            para = next_para()
            write_runs(para, block.runs)
            style_subheading(para, block.level)
        elif isinstance(block, CodeBlock):
            _fill_body_code(next_para, block, width_pt, avail_in)
        elif isinstance(block, ListBlock):
            _fill_body_list(next_para, block)
        elif isinstance(block, BlockQuote):
            _fill_body_quote(next_para, block)


def _fill_body_code(
    next_para: Callable[[], _Paragraph],
    block: CodeBlock,
    width_pt: float,
    avail_in: float,
) -> None:
    # inside a column there is no room for the bordered flow box: code
    # renders as monospace paragraphs in the column text box, font-fitted
    # to the column width/band (wrap is native)
    font_pt, _height = fit_code(block.text, width_pt, avail_in)
    for _idx, code_line in enumerate(block.text.splitlines() or [""]):
        para = next_para()
        run = para.add_run()
        run.text = code_line or " "
        run.font.name = MONO_FONT
        run.font.size = Pt(font_pt)


def _fill_body_list(next_para: Callable[[], _Paragraph], block: ListBlock) -> None:
    for idx, item in enumerate(block.items):
        para = next_para()
        write_runs(para, item.runs)
        apply_list_para(
            para,
            block.ordered,
            item.level,
            start=block.start if idx == 0 else None,
        )


def _fill_body_quote(next_para: Callable[[], _Paragraph], block: BlockQuote) -> None:
    for quoted in block.paragraphs:
        para = next_para()
        write_runs(para, quote_runs(quoted))
        apply_quote_para(para)


def estimate_height(block: Block, images: Mapping[int, bytes]) -> int:
    if isinstance(block, Paragraph):
        return int(PARA_HEIGHT)
    if isinstance(block, Subheading):
        return int(SUBHEAD_HEIGHT)
    if isinstance(block, CodeBlock):
        _font, height = fit_code(block.text, CONTENT_WIDTH_PT, FLOW_BAND_IN)
        return int(Inches(height))
    if isinstance(block, ListBlock):
        return int(LIST_HEIGHT) * len(block.items)
    if isinstance(block, BlockQuote):
        return int(QUOTE_HEIGHT) * max(1, len(block.paragraphs))
    if isinstance(block, TableBlock):
        return int(TABLE_ROW_HEIGHT) * len(block.rows) + int(Inches(0.15))
    if isinstance(block, ImageBlock):
        return picture_geometry(images[block.index])[1]
    return int(PARA_HEIGHT)


def flow_blocks(
    slide: Slide,
    blocks: Iterable[Block],
    bottom: int = FLOW_BOTTOM,
    images: Mapping[int, bytes] | None = None,
) -> None:
    """Render blocks as sequential native shapes down the content band,
    each fitted into the space actually left."""
    images = images or {}
    y = int(FLOW_TOP)
    for block in blocks:
        avail_in = max(0.3, (int(bottom) - y) / 914400)
        if isinstance(block, CodeBlock):
            # fitted to the space actually left: font shrinks, wrap is on,
            # the box height matches the wrapped content
            font_pt, height_in = fit_code(block.text, CONTENT_WIDTH_PT, avail_in)
            height = min(int(Inches(height_in)), int(bottom) - y) or int(Inches(0.3))
            add_code_box(slide, y, height, block, font_pt)
        else:
            height = _flow_non_code(slide, block, y, int(bottom), images)
        y += height + int(BLOCK_GAP)


def _flow_non_code(
    slide: Slide, block: Block, y: int, bottom: int, images: Mapping[int, bytes]
) -> int:
    height = min(estimate_height(block, images), bottom - y) if y < bottom else 0
    if height <= 0:
        height = int(Inches(0.3))  # overflow: keep a minimal, clamped shape
    if isinstance(block, TableBlock):
        add_table(slide, y, height, block)
    elif isinstance(block, ImageBlock):
        add_picture(slide, y, block, images[block.index])
    elif isinstance(block, Paragraph):
        add_text_block(slide, y, height, [block.runs])
    elif isinstance(block, Subheading):
        paras = add_text_block(slide, y, height, [block.runs])
        style_subheading(paras[0], block.level)
    elif isinstance(block, BlockQuote):
        add_text_block(slide, y, height, [quote_runs(p) for p in block.paragraphs])
    elif isinstance(block, ListBlock):
        _flow_list(slide, y, height, block)
    return height


def _flow_list(slide: Slide, y: int, height: int, block: ListBlock) -> None:
    paragraphs = [item.runs for item in block.items]
    box_paragraphs = add_text_block(slide, y, height, paragraphs)
    for idx, (para, item) in enumerate(zip(box_paragraphs, block.items, strict=False)):
        apply_list_para(
            para,
            block.ordered,
            item.level,
            start=block.start if idx == 0 else None,
        )
