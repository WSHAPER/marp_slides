"""Native content elements: runs, paragraphs and shapes.

Paragraph-level styling (run properties, native list and quote formatting,
subheading sizing) and the shape adders the flow layout composes slides from
(code boxes, tables, text blocks, pictures, chrome boxes, the slide-number
field). Everything here produces ordinary, editable PowerPoint objects.
"""

from __future__ import annotations

import io
from collections.abc import Iterable

from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.parts.image import Image as PPTXImage
from pptx.slide import Slide
from pptx.text.text import _Paragraph
from pptx.util import Emu, Pt

from .blocks import (
    CodeBlock,
    ImageBlock,
    Paragraph,
    Run,
    TableBlock,
)
from .geometry import (
    CHROME_BOX_HEIGHT,
    CHROME_COLOR,
    CHROME_FONT_SIZE,
    CODE_BG,
    CODE_BORDER,
    CONTENT_LEFT,
    CONTENT_WIDTH,
    EMU_PER_INCH,
    FOOTER_TOP,
    LIST_MARL,
    MONO_FONT,
    PICTURE_MAX_HEIGHT,
    SLIDENUM_FIELD_GUID,
    SLIDENUM_LEFT,
    SLIDENUM_WIDTH,
    SUBHEAD_SIZES,
)


def append_runs(para: _Paragraph, runs: Iterable[Run]) -> None:
    for run in runs:
        r = para.add_run()
        r.text = run.text
        r.font.bold = run.bold or None
        r.font.italic = run.italic or None
        if run.code:
            r.font.name = MONO_FONT
        if run.link:
            r.hyperlink.address = run.link  # native text-run hyperlink


def write_runs(para: _Paragraph, runs: Iterable[Run]) -> None:
    para.clear()  # type: ignore[no-untyped-call]
    append_runs(para, runs)


def style_subheading(para: _Paragraph, level: int) -> None:
    """Subheadings render as bold text larger than body, smaller than the
    slide title (stepped by heading level, clamped at the body size)."""
    para.space_before = Pt(8)
    for r in para.runs:
        r.font.bold = True
        r.font.size = SUBHEAD_SIZES[min(max(level, 3), 6) - 3]


def apply_list_para(
    para: _Paragraph, ordered: bool, level: int, start: int | None = None
) -> None:
    """Native list formatting. ``start`` (an ordered run's CommonMark start
    number) is emitted as buAutoNum ``startAt``: only the FIRST paragraph
    of a run may carry it, or every item would restart the numbering."""
    pPr = para._p.get_or_add_pPr()
    # native outline level (0-8), plus hanging-indent list formatting
    pPr.set("lvl", str(min(8, max(0, level))))
    pPr.set("marL", str(int(LIST_MARL) * (level + 1)))
    pPr.set("indent", f"-{LIST_MARL}")
    if ordered:
        attrs = {"type": "arabicPeriod"}
        if start is not None and start != 1:
            attrs["startAt"] = str(start)
        bu = pPr.makeelement(qn("a:buAutoNum"), attrs)
    else:
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def apply_quote_para(para: _Paragraph) -> None:
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", LIST_MARL)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def quote_runs(paragraph: Paragraph) -> tuple[Run, ...]:
    """A quote's runs carry the same styling as the source, forced italic."""
    return tuple(
        Run(
            r.text,
            bold=r.bold,
            italic=True,
            code=r.code,
            link=r.link,
        )
        for r in paragraph.runs
    )


def picture_geometry(blob: bytes) -> tuple[int, int]:
    """Fitted (width, height) in EMU for a picture blob: native size from the
    image's own pixel dimensions and dpi metadata, scaled down (never up) to
    fit the content width and the picture height cap, aspect preserved."""
    image = PPTXImage.from_blob(blob)
    w_px, h_px = image.size
    dpi_x, dpi_y = image.dpi
    native_w = int(w_px * EMU_PER_INCH / max(1, dpi_x))
    native_h = int(h_px * EMU_PER_INCH / max(1, dpi_y))
    scale = min(
        1.0,
        int(CONTENT_WIDTH) / native_w if native_w else 1.0,
        int(PICTURE_MAX_HEIGHT) / native_h if native_h else 1.0,
    )
    return int(native_w * scale), int(native_h * scale)


def add_picture(slide: Slide, y: int, block: ImageBlock, blob: bytes) -> None:
    """Insert a file-id image as a native, editable picture shape: fitted
    geometry, horizontally centered, alt text as the native description."""
    width, height = picture_geometry(blob)
    left = int(CONTENT_LEFT) + (int(CONTENT_WIDTH) - width) // 2
    pic = slide.shapes.add_picture(
        io.BytesIO(blob), Emu(left), Emu(y), Emu(width), Emu(height)
    )
    if block.alt:
        # pptx oxml typing gap: _element is only partially typed upstream
        pic._element.nvPicPr.cNvPr.set("descr", block.alt)  # type: ignore[union-attr]


def add_code_box(
    slide: Slide, y: int, height: int, block: CodeBlock, font_pt: float
) -> None:
    """Bordered monospace code box. word_wrap is ON: every renderer wraps at
    the box width (deterministic; see the code-geometry notes in geometry),
    and the height/font were fitted by the code-fit ladder so the wrapped
    content stays inside the box."""
    box = slide.shapes.add_textbox(CONTENT_LEFT, Emu(y), CONTENT_WIDTH, Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, line in enumerate(block.text.splitlines() or [""]):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        r = para.add_run()
        r.text = line or " "
        r.font.name = MONO_FONT
        r.font.size = Pt(font_pt)
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = CODE_BORDER
    box.line.width = Pt(0.75)


def add_table(slide: Slide, y: int, height: int, block: TableBlock) -> None:
    rows = len(block.rows)
    cols = len(block.rows[0]) if block.rows else 1
    graphic = slide.shapes.add_table(
        rows, cols, CONTENT_LEFT, Emu(y), CONTENT_WIDTH, Emu(height)
    )
    table = graphic.table
    for col in range(cols):
        table.columns[col].width = Emu(int(CONTENT_WIDTH) // cols)
    for row_idx, row in enumerate(block.rows):
        for col_idx, cell_runs in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            para = cell.text_frame.paragraphs[0]
            append_runs(para, cell_runs)  # rich inline runs, not flattened
            if row_idx == 0:  # header row stays bold
                for run in para.runs:
                    run.font.bold = True


def add_text_block(
    slide: Slide, y: int, height: int, paragraphs: Iterable[Iterable[Run]]
) -> list[_Paragraph]:
    """Add a plain textbox with one paragraph per run-list; returns the
    paragraph objects (so callers can apply list formatting)."""
    box = slide.shapes.add_textbox(CONTENT_LEFT, Emu(y), CONTENT_WIDTH, Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    created = []
    for idx, runs in enumerate(paragraphs):
        para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        created.append(para)
        append_runs(para, runs)
    return created


def add_chrome_box(slide: Slide, runs: Iterable[Run], top: int, name: str) -> None:
    """Deck-level footer/header: a small, gray, centered text box on every
    slide, outside the content flow (native and editable like everything
    else: it is an ordinary textbox the user can restyle)."""
    box = slide.shapes.add_textbox(
        CONTENT_LEFT, Emu(int(top)), CONTENT_WIDTH, Emu(int(CHROME_BOX_HEIGHT))
    )
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    append_runs(para, runs)
    para.alignment = PP_ALIGN.CENTER
    for r in para.runs:
        r.font.size = CHROME_FONT_SIZE
        r.font.color.rgb = CHROME_COLOR


def add_slide_number_field(slide: Slide, number: int) -> None:
    """Native slide number (frontmatter 'paginate: true'): a real ``<a:fld
    type="slidenum">`` field in a small bottom-right text box, NOT literal
    text, so the number updates automatically when slides are added,
    removed or reordered."""
    box = slide.shapes.add_textbox(
        SLIDENUM_LEFT, Emu(int(FOOTER_TOP)), SLIDENUM_WIDTH, Emu(int(CHROME_BOX_HEIGHT))
    )
    box.name = "Slide Number"
    tf = box.text_frame
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    fld = para._p.makeelement(qn("a:fld"), {})
    fld.set("id", SLIDENUM_FIELD_GUID)
    fld.set("type", "slidenum")
    rpr = fld.makeelement(qn("a:rPr"), {})
    rpr.set("lang", "en-US")
    rpr.set("sz", "1000")
    fill = rpr.makeelement(qn("a:solidFill"), {})
    clr = fill.makeelement(qn("a:srgbClr"), {"val": "7F7F7F"})
    fill.append(clr)
    rpr.append(fill)
    fld.append(rpr)
    t = fld.makeelement(qn("a:t"), {})
    t.text = str(number)  # naive-renderer fallback; field-aware apps re-derive it
    fld.append(t)
    para._p.append(fld)
